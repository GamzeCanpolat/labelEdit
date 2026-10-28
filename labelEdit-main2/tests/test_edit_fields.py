"""İşlem sayfası alanlarının kayıt ve dışa aktarma sözleşmeleri."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from conftest import maps_url, write_xlsx
from edit_data_manager import EditDataManager, export_excel
from edit_fields import field_columns


def source(input_dir, name, with_fields=True):
    headers = ["KONTROL", "isim", "maps_url"]
    row = ["İncele", "Mağaza", maps_url(1)]
    if with_fields:
        headers += ["brand_name", "mainCategory", "Alt Kategori", "Pol State"]
        row += ["Eski Marka", "Alışveriş", "Mobilya", "CLOSED_TEMPORARILY"]
    return write_xlsx(input_dir / name, {"Mekanlar": [headers, row]})


@pytest.mark.parametrize("with_fields", [True, False])
def test_edit_fields_persist_and_export(input_dir, tmp_path, with_fields):
    p = source(input_dir, f"cleanpoi_cevre_fields_{with_fields}.xlsx", with_fields)
    dm = EditDataManager(p.name)
    row = dm.get_place_records()[0].row_index
    assert dm.get_edit_fields(row)["brand"] == ("Eski Marka" if with_fields else "")
    fields = {"brand": "Yeni Marka", "main_category": "Ev", "sub_category": "Dekorasyon"}
    dm.keep_row(row, fields)
    dm.undo_action(row)
    assert dm.get_edit_fields(row) == {**fields, "poi_state": "CLOSED_TEMPORARILY" if with_fields else ""}
    dm.close()
    dm = EditDataManager(p.name)
    assert dm.get_edit_fields(row)["brand"] == "Yeni Marka"
    dm.close()
    out = tmp_path / "export.xlsx"
    export_excel(out, p.name)
    sheet = openpyxl.load_workbook(out).active
    columns = [c.value for c in sheet[1]]
    mapping = field_columns(columns)
    assert sheet.cell(2, columns.index(mapping["brand"]) + 1).value == "Yeni Marka"
    assert sheet.cell(2, columns.index(mapping["sub_category"]) + 1).value == "Dekorasyon"
    if with_fields:
        assert not ({"Marka", "Ana Kategori"} & set(columns))


def test_workspace_edit_api_keeps_fields_and_preserves_poi_state(input_dir, tmp_path, monkeypatch):
    import main
    from workspace_manager import workspace_mgr

    p = source(input_dir, "workspace_edit_fields.xlsx")
    ws = workspace_mgr.create_workspace(p, "İşlem alanları")
    ws_id = ws["id"]
    workspace_mgr.update_column_mapping(ws_id, {"status": "Pol State"})
    monkeypatch.setattr(main.state, "edit_workspace_id", ws_id)
    # Etiketleme sayfasının aktif ortamı farklı olsa da işlem kaydı doğru yere gider.
    monkeypatch.setattr(main.state, "active_workspace_id", "another-workspace")
    with TestClient(main.app) as client:
        options = client.get("/api/edit/field-options")
        assert options.json()["brand"] == ["Eski Marka"]
        fields = {"brand": 'Marka "A"', "main_category": "Ev", "sub_category": "Dekorasyon"}
        response = client.post("/api/edit/action", json={"row_index": 0, "action": "keep", "fields": fields})
        assert response.status_code == 200, response.text
        assert workspace_mgr.get_edit_fields(ws_id, 0) == {**fields, "poi_state": "CLOSED_TEMPORARILY"}
        assert client.post("/api/edit/action/undo", json={"row_index": 0}).status_code == 200
        assert client.post("/api/edit/action", json={"row_index": 0, "action": "skip", "fields": {"brand": "Kaydedilmemeli"}}).status_code == 200
        assert workspace_mgr.get_edit_fields(ws_id, 0)["brand"] == 'Marka "A"'
        assert client.post("/api/edit/action", json={"row_index": 0, "action": "keep", "fields": {"_status": "silindi"}}).status_code == 422
        assert client.post("/api/edit/action", json={"row_index": 999, "action": "keep", "fields": fields}).status_code == 404
        result = client.get("/api/edit/export")
        assert result.status_code == 200
        sheet = openpyxl.load_workbook(io.BytesIO(result.content)).active
        columns = [c.value for c in sheet[1]]
        assert sheet.cell(2, columns.index("brand_name") + 1).value == 'Marka "A"'


def test_workspace_adds_missing_fields(input_dir):
    from workspace_manager import workspace_mgr
    p = source(input_dir, "workspace_without_fields.xlsx", False)
    ws = workspace_mgr.create_workspace(p, "Eksik alanlar")
    workspace_mgr.update_action(ws["id"], 0, "keep", {"brand": "Yeni", "sub_category": "Alt"})
    assert workspace_mgr.get_edit_fields(ws["id"], 0)["brand"] == "Yeni"
    assert {"Marka", "Alt Kategori"} <= set(workspace_mgr.get_workspace(ws["id"])["columns"])


def test_workspace_decisions_survive_new_manager_without_rewriting_source(input_dir):
    from workspace_manager import WorkspaceManager, workspace_mgr
    path = write_xlsx(input_dir / "persistent_workspace.xlsx", {"Veriler": [
        ("name", "maps", "Marka", "Ana Kategori", "Alt Kategori"),
        ("A", maps_url(1), "Eski", "Ev", "Mobilya"),
        ("B", maps_url(2), "Eski", "Ev", "Mobilya"),
        ("C", maps_url(3), "Eski", "Ev", "Mobilya"),
    ]})
    before = path.read_bytes()
    ws = workspace_mgr.create_workspace(path, "Kalıcı kayıt")
    workspace_mgr.update_action(ws["id"], 0, "keep", {"brand": "Yeni Marka", "main_category": "Alışveriş", "sub_category": "Dekorasyon"})
    workspace_mgr.update_action(ws["id"], 1, "delete")
    workspace_mgr.update_action(ws["id"], 2, "skip", {"brand": "Kaydedilmemeli"})
    reopened = WorkspaceManager(workspace_mgr.db_path)
    assert reopened.get_edit_fields(ws["id"], 0)["brand"] == "Yeni Marka"
    assert reopened.get_edit_fields(ws["id"], 0)["sub_category"] == "Dekorasyon"
    assert reopened.get_edit_fields(ws["id"], 2)["brand"] == "Eski"
    assert reopened.get_workspace_stats(ws["id"])["kept"] == 1
    assert reopened.get_workspace_stats(ws["id"])["deleted"] == 1
    assert [r.row_index for r in reopened.get_place_records(ws["id"], status_filter="bekliyor")] == [2]
    assert path.read_bytes() == before  # İşlemler SQLite'a kaydedildi; kaynak Excel değişmedi.


@pytest.mark.parametrize('with_column', [False, True])
@pytest.mark.parametrize('state', ['OPERATIONAL', 'CLOSED_PERMANENTLY', 'CLOSED_TEMPORARILY'])
def test_poi_state_is_saved_and_exported_in_both_storage_flows(input_dir, tmp_path, with_column, state):
    from workspace_manager import WorkspaceManager, workspace_mgr
    p = source(input_dir, f'cleanpoi_cevre_state_{with_column}_{state}.xlsx', with_column)
    if with_column:
        wb = openpyxl.load_workbook(p)
        wb.active.cell(2, wb.active.max_column).value = None
        wb.save(p)
        wb.close()
    manager = EditDataManager(p.name)
    manager.keep_row(0, {'poi_state': state})
    manager.close()
    manager = EditDataManager(p.name)
    assert manager.get_edit_fields(0)['poi_state'] == state
    manager.close()
    out = tmp_path / 'legacy.xlsx'
    export_excel(out, p.name)
    sheet = openpyxl.load_workbook(out).active
    columns = [c.value for c in sheet[1]]
    column = field_columns(columns)['poi_state']
    assert sheet.cell(2, columns.index(column) + 1).value == state

    ws = workspace_mgr.create_workspace(p, 'Durum seçimi')
    old = workspace_mgr.get_edit_fields(ws['id'], 0)['poi_state']
    workspace_mgr.update_action(ws['id'], 0, 'skip', {'poi_state': state})
    assert workspace_mgr.get_edit_fields(ws['id'], 0)['poi_state'] == old
    workspace_mgr.update_action(ws['id'], 0, 'keep', {'poi_state': state})
    reopened = WorkspaceManager(workspace_mgr.db_path)
    assert reopened.get_edit_fields(ws['id'], 0)['poi_state'] == state
    out = tmp_path / 'workspace.xlsx'
    reopened.export_to_excel(ws['id'], out)
    sheet = openpyxl.load_workbook(out).active
    columns = [c.value for c in sheet[1]]
    column = field_columns(columns)['poi_state']
    assert sheet.cell(2, columns.index(column) + 1).value == state


def test_invalid_poi_state_cannot_partially_save_other_fields(input_dir, monkeypatch):
    import main
    from workspace_manager import workspace_mgr
    p = source(input_dir, 'invalid_poi_state.xlsx')
    ws = workspace_mgr.create_workspace(p, 'Durum doğrulama')
    monkeypatch.setattr(main.state, 'edit_workspace_id', ws['id'])
    with TestClient(main.app) as client:
        result = client.post('/api/edit/action', json={
            'row_index': 0, 'action': 'keep', 'fields': {'brand': 'Kaydedilmemeli', 'poi_state': 'INVALID_STATUS'}})
        assert result.status_code == 422
        assert workspace_mgr.get_edit_fields(ws['id'], 0)['brand'] == 'Eski Marka'


def test_existing_poi_state_is_protected_in_both_storage_flows(input_dir, monkeypatch):
    import main
    from workspace_manager import workspace_mgr
    p = source(input_dir, 'cleanpoi_cevre_known_state.xlsx')
    manager = EditDataManager(p.name)
    for value in ['OPERATIONAL', 'CLOSED_PERMANENTLY', '']:
        with pytest.raises(ValueError, match='Mevcut Pol State'):
            manager.keep_row(0, {'brand': 'Kaydedilmemeli', 'poi_state': value})
        assert manager.get_edit_fields(0)['poi_state'] == 'CLOSED_TEMPORARILY'
        assert manager.get_edit_fields(0)['brand'] == 'Eski Marka'
    manager.close()
    ws = workspace_mgr.create_workspace(p, 'Mevcut durum korunur')
    monkeypatch.setattr(main.state, 'edit_workspace_id', ws['id'])
    with TestClient(main.app) as client:
        result = client.post('/api/edit/action', json={
            'row_index': 0, 'action': 'keep', 'fields': {'brand': 'Kaydedilmemeli', 'poi_state': 'OPERATIONAL'}})
        assert result.status_code == 422
        assert workspace_mgr.get_edit_fields(ws['id'], 0)['poi_state'] == 'CLOSED_TEMPORARILY'
        assert workspace_mgr.get_edit_fields(ws['id'], 0)['brand'] == 'Eski Marka'
        result = client.post('/api/edit/action', json={
            'row_index': 0, 'action': 'keep', 'fields': {'brand': 'Yeni Marka'}})
        assert result.status_code == 200
        assert workspace_mgr.get_edit_fields(ws['id'], 0)['poi_state'] == 'CLOSED_TEMPORARILY'
