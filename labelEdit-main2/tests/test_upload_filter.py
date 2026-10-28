"""Büyük dosya metadatası tekrar okunmamalı; kişi filtresi diğer satırları korumalı."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from conftest import maps_url, write_xlsx
from workspace_manager import workspace_mgr


def assigned_file(input_dir, name='assigned.xlsx'):
    return write_xlsx(input_dir / name, {'Veri': [
        ('İsim', 'Maps_Link', 'Kontrol_Eden'),
        ('G1', maps_url(1), 'Gamze'),
        ('A1', maps_url(2), 'Abdullah'),
        ('G2', None, ' Gamze '),
        ('X1', maps_url(4), None),
    ], 'Listeler': [('Kişi',), ('Gamze',), ('Abdullah',)]})


def test_workspace_details_do_not_reopen_large_excel(input_dir, monkeypatch):
    import workspace_manager
    path = assigned_file(input_dir, 'metadata.xlsx')
    ws = workspace_mgr.create_workspace(path, 'Hızlı metadata')
    def forbidden(*args, **kwargs):
        raise AssertionError('Liste/detay için Excel satırları yeniden okunmamalı.')
    monkeypatch.setattr(workspace_manager.openpyxl, 'load_workbook', forbidden)
    info = workspace_mgr.get_workspace(ws['id'])
    assert info['available_sheets'] == ['Veri', 'Listeler']
    assert info['stats']['total'] == 4
    assert any(w['id'] == ws['id'] for w in workspace_mgr.list_workspaces())


def test_assignee_counts_filtering_and_sql_parameters(input_dir):
    path = assigned_file(input_dir)
    ws = workspace_mgr.create_workspace(path, 'Atamalar')
    options = workspace_mgr.get_assignee_options(ws['id'])
    assert options == {'column': 'Kontrol_Eden', 'choices': [
        {'value': 'Abdullah', 'total': 1, 'pending': 1},
        {'value': 'Gamze', 'total': 2, 'pending': 2},
    ]}
    rows = workspace_mgr.get_place_records(ws['id'], status_filter='bekliyor', assignee='Gamze')
    assert [r.row_index for r in rows] == [0, 2]
    assert [r.position for r in rows] == [0, 1] and all(r.total == 2 for r in rows)
    assert workspace_mgr.get_place_records(ws['id'], assignee="Gamze' OR 1=1 --") == []
    workspace_mgr.update_action(ws['id'], 0, 'keep')
    assert [r.row_index for r in workspace_mgr.get_place_records(ws['id'], status_filter='bekliyor', assignee='Gamze')] == [2]
    assert workspace_mgr.get_assignee_options(ws['id'])['choices'][1]['pending'] == 1
    assert workspace_mgr.get_assignee_pending_count(ws['id'], 'Gamze') == 1


def test_filtered_session_cannot_change_another_persons_row(input_dir, monkeypatch):
    import main
    path = assigned_file(input_dir, 'filtered_session.xlsx')
    ws = workspace_mgr.create_workspace(path, 'Gamze oturumu')
    class Session:
        def start(self): pass
        def stop(self): pass
    class Worker:
        def __init__(self, session): self.records = []
        async def start(self, records): self.records = records
        async def stop(self): pass
    monkeypatch.setattr(main, 'SessionManager', Session)
    monkeypatch.setattr(main, 'ScreenshotWorker', Worker)
    with TestClient(main.app) as client:
        schema = client.get('/api/edit/schema', params={'workspace_id': ws['id']}).json()
        assert schema['assignees']['column'] == 'Kontrol_Eden'
        start = client.post('/api/edit/session/start', json={'workspace_id': ws['id'], 'assignee': 'Gamze'})
        assert start.status_code == 200, start.text
        assert start.json()['pending'] == 2 and start.json()['assignee'] == 'Gamze'
        assert start.json()['total'] == 2  # Atanmış fakat bağlantısı boş kayıt da kaybolmaz.
        assert client.post('/api/edit/action', json={'row_index': 1, 'action': 'delete'}).status_code == 422
        assert client.post('/api/edit/action/undo', json={'row_index': 1}).status_code == 422
        assert client.post('/api/edit/action', json={'row_index': 0, 'action': 'keep'}).status_code == 200
        assert workspace_mgr.get_workspace_stats(ws['id'])['deleted'] == 0


def test_upload_keeps_original_filename(input_dir):
    import main
    wb = openpyxl.Workbook()
    wb.active.append(['İsim', 'Maps_Link'])
    wb.active.append(['Yer', maps_url(1)])
    data = io.BytesIO()
    wb.save(data)
    with TestClient(main.app) as client:
        response = client.post('/api/workspaces/upload', files={
            'file': ('POI_Birlesik_Veri.xlsx', data.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        assert response.status_code == 200, response.text
        assert response.json()['source_file'] == 'POI_Birlesik_Veri.xlsx'
        assert response.json()['stored_file_path'].endswith('POI_Birlesik_Veri.xlsx')
