"""
Çalışma Ortamları (Workspaces), Dinamik Sütun Eşleştirme ve SQLite Entegrasyon Testleri
"""
import io
import openpyxl
from fastapi.testclient import TestClient
import main
from workspace_manager import workspace_mgr


def test_workspace_workflow_end_to_end():
    # 1. Bellekte sentetik bir Excel oluştur
    # Sütun isimleri bilerek farklı: "Mekan Bağlantısı", "Firma Başlığı", "Mevcut Durum"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Firmalar"
    ws.append(["Firma Başlığı", "Mekan Bağlantısı", "Mevcut Durum"])
    for i in range(5):
        ws.append([f"Firma {i}", f"https://www.google.com/maps/place/Firma{i}", "bekliyor"])

    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)

    with TestClient(main.app) as client:
        # 2. Dosyayı yükle
        response = client.post(
            "/api/workspaces/upload",
            files={"file": ("test_harita_verisi.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"name": "Test Mobilya Ortamı"},
        )
        assert response.status_code == 200, response.text
        data = response.json()
        ws_id = data["id"]
        assert data["name"] == "Test Mobilya Ortamı"
        assert "Mekan Bağlantısı" in data["columns"]
        assert "Firma Başlığı" in data["columns"]

        # 3. Çalışma ortamlarını listele
        list_res = client.get("/api/workspaces")
        assert list_res.status_code == 200
        workspaces = list_res.json()
        assert any(w["id"] == ws_id for w in workspaces)

        # 4. Sütun ekleme (Kategori sütunu Excel'de yok, SQLite tablosuna dinamik ekliyoruz)
        col_res = client.post(
            f"/api/workspaces/{ws_id}/columns/add",
            json={"name": "Kategori", "default_value": ""},
        )
        assert col_res.status_code == 200
        updated_cols = col_res.json()["workspace"]["columns"]
        assert "Kategori" in updated_cols

        # İkinci bir ek sütun ekle: "İnceleme Notu"
        col_res2 = client.post(
            f"/api/workspaces/{ws_id}/columns/add",
            json={"name": "İnceleme Notu", "default_value": ""},
        )
        assert col_res2.status_code == 200
        assert "İnceleme Notu" in col_res2.json()["workspace"]["columns"]

        # 5. Dinamik Sütun Eşleştirme (Mapping)
        # "Mekan Bağlantısı" -> maps_url
        # "Firma Başlığı" -> name
        # "Kategori" -> category
        # "Mevcut Durum" -> status
        map_res = client.post(
            f"/api/workspaces/{ws_id}/mapping",
            json={
                "mapping": {
                    "maps_url": "Mekan Bağlantısı",
                    "name": "Firma Başlığı",
                    "category": "Kategori",
                    "status": "Mevcut Durum",
                }
            },
        )
        assert map_res.status_code == 200
        mapping = map_res.json()["workspace"]["column_mapping"]
        assert mapping["maps_url"] == "Mekan Bağlantısı"
        assert mapping["name"] == "Firma Başlığı"
        assert mapping["category"] == "Kategori"

        # 6. Aktif workspace yap
        act_res = client.post(f"/api/workspaces/{ws_id}/activate")
        assert act_res.status_code == 200
        assert act_res.json()["status"] == "activated"

        # 7. Etiketleme oturumu başlat (workspace_id ile)
        start_res = client.post("/api/session/start", json={"start_index": 0, "workspace_id": ws_id})
        assert start_res.status_code == 200, start_res.text
        start_data = start_res.json()
        assert start_data["status"] == "started"
        assert start_data["total"] == 5

        # 8. SQLite tablosuna doğrudan etiket kaydet (Kategori ve ekstra sütun güncellemesi)
        label_res = client.post(
            "/api/label",
            json={
                "row_index": 0,
                "category": "Mobilya & Dekorasyon",
                "extra_fields": {"İnceleme Notu": "Onaylandı"},
            },
        )
        assert label_res.status_code == 200, label_res.text
        assert label_res.json()["status"] == "saved"

        # Oturumu durdur
        client.post("/api/session/stop")

        # 9. Düzenle akışı (Düzenle oturumu başlat ve aksiyon ver)
        edit_start_res = client.post("/api/edit/session/start", json={"start_index": 0, "workspace_id": ws_id})
        assert edit_start_res.status_code == 200, edit_start_res.text

        # Satır 1 için 'delete' aksiyonu ver
        action_res = client.post("/api/edit/action", json={"row_index": 1, "action": "delete"})
        assert action_res.status_code == 200
        assert action_res.json()["status"] == "deleted"

        # Satır 2 için 'keep' aksiyonu ver
        keep_res = client.post("/api/edit/action", json={"row_index": 2, "action": "keep"})
        assert keep_res.status_code == 200
        assert keep_res.json()["status"] == "kept"

        # Düzenle oturumunu durdur
        client.post("/api/edit/session/stop")

        # 10. Excel Export (SQLite tablosundaki güncel halin indirilmesi)
        export_res = client.get(f"/api/workspaces/{ws_id}/export")
        assert export_res.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in export_res.headers["content-type"]

        # İndirilen Excel'in içeriğini doğrula
        exported_wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
        sheet = exported_wb.active
        headers = [c.value for c in sheet[1]]
        assert "Mekan Bağlantısı" in headers
        assert "Firma Başlığı" in headers
        assert "Kategori" in headers
        assert "İnceleme Notu" in headers

        # Kaydedilen kategori değerini doğrula
        cat_col_idx = headers.index("Kategori") + 1
        note_col_idx = headers.index("İnceleme Notu") + 1
        assert sheet.cell(row=2, column=cat_col_idx).value == "Mobilya & Dekorasyon"
        assert sheet.cell(row=2, column=note_col_idx).value == "Onaylandı"

        # 11. Temizlik: Workspace'i sil
        del_res = client.delete(f"/api/workspaces/{ws_id}")
        assert del_res.status_code == 200
